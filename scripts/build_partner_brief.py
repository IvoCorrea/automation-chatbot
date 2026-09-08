from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "Proposta_MVP_Agente_Imobiliario.docx"
NAVY = "0B2545"; BLUE = "2E74B5"; DARK = "1F4D78"; MUTED = "5B677A"; LIGHT = "F2F4F7"; PALE = "E8EEF5"; GOLD = "B7791F"; WHITE = "FFFFFF"

def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar = OxmlElement("w:tcMar"); tcPr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None: node = OxmlElement(f"w:{side}"); mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")

def set_cell_text(cell, text, bold=False, color=NAVY, size=10.2):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text); set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margins(cell)

def fixed_table(table, widths):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr; layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    table.width = Inches(6.5)
    for row in table.rows:
        for cell, width in zip(row.cells, widths): cell.width = Inches(width)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p

def add_body(doc, text, emphasis=None):
    p = doc.add_paragraph(style="Normal")
    if emphasis and emphasis in text:
        before, after = text.split(emphasis, 1)
        p.add_run(before); r = p.add_run(emphasis); set_font(r, size=11, color=NAVY, bold=True); p.add_run(after)
    else: p.add_run(text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.9); section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1); section.right_margin = Inches(1)
section.header_distance = Inches(0.49); section.footer_distance = Inches(0.49)

styles = doc.styles
normal = styles["Normal"]; normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri"); normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.28
for name, size, color, before, after in [("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,12,6),("Heading 3",12,DARK,8,4)]:
    s = styles[name]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); s._element.rPr.rFonts.set(qn("w:hAnsi"),"Calibri"); s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color); s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)

# Header and footer.
header_p = section.header.paragraphs[0]; header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header_p.add_run("IMOB AGENT | VISÃO PARA SÓCIOS"); set_font(r, size=8.5, color=MUTED, bold=True)
footer_p = section.footer.paragraphs[0]; footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer_p.add_run("Documento de trabalho | Agosto de 2026"); set_font(r, size=8.5, color=MUTED)

# Cover: proposal_centerpiece header pattern.
doc.add_paragraph().paragraph_format.space_after = Pt(55)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("IMOB AGENT"); set_font(r, size=13, color=GOLD, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
r = p.add_run("Atendimento inteligente no WhatsApp\npara transformar leads em visitas"); set_font(r, size=27, color=NAVY, bold=True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(28)
r = p.add_run("Proposta de MVP e estratégia comercial para o setor imobiliário"); set_font(r, size=14, color=MUTED)

t = doc.add_table(rows=2, cols=2); fixed_table(t, [3.25, 3.25])
cover = [("Tese", "Responder em segundos, qualificar melhor e entregar o lead certo ao corretor."), ("Primeiro nicho", "Imobiliárias que recebem leads por WhatsApp e anúncios digitais."), ("Produto inicial", "Agente de IA + busca de imóveis + handoff humano."), ("Objetivo", "Validar receita recorrente com pilotos antes de generalizar para outros setores.")]
for cell, (label, value) in zip([c for row in t.rows for c in row.cells], cover):
    shade(cell, PALE); cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); r=p.add_run(label.upper()); set_font(r,size=8.5,color=GOLD,bold=True)
    p=cell.add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run(value); set_font(r,size=10.2,color=NAVY,bold=True); cell_margins(cell,150,180,150,180)

doc.add_paragraph().paragraph_format.space_after = Pt(35)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("A oportunidade não é substituir o corretor. É garantir que ele receba um cliente preparado para avançar."); set_font(r,size=12,color=DARK,italic=True)
doc.add_page_break()

add_heading(doc, "1. A visão: criar velocidade com qualidade", 1)
add_body(doc, "O WhatsApp já é a porta de entrada da maior parte das conversas comerciais no Brasil. Em imobiliárias, o problema não é a falta de leads; é a perda de intenção entre o anúncio e o primeiro atendimento humano. Quando a resposta demora, o interessado esfria, conversa com outro corretor ou simplesmente desaparece.")
add_body(doc, "Nossa proposta é simples: um agente que responde imediatamente, entende o que a pessoa procura, consulta o catálogo disponível e faz a passagem para um corretor no momento certo. Isso cria uma operação mais rápida para o cliente e uma experiência mais fluida para o comprador ou locatário.")
add_heading(doc, "O que vendemos", 2)
add_bullets(doc, [
    "Atendimento 24/7 no canal onde o lead já está: WhatsApp.",
    "Qualificação consistente: finalidade, região, faixa de preço, tipo de imóvel, prazo e financiamento.",
    "Sugestão de imóveis compatíveis, sem prometer disponibilidade que não exista.",
    "Transferência organizada para o corretor, com histórico e contexto da conversa.",
    "Uma base que pode ser configurada para outros setores após validação no imobiliário."
])
add_heading(doc, "Por que começar em imobiliárias", 2)
add_body(doc, "É um nicho com dor clara, tickets relevantes, jornadas consultivas e informação estruturada - catálogo, preço, região e características. Também permite medir valor rapidamente: tempo de primeira resposta, taxa de qualificação, visitas agendadas e conversão de lead para oportunidade.")

add_heading(doc, "2. Como a aplicação funciona, passo a passo", 1)
add_body(doc, "A operação foi desenhada para ser útil desde a primeira mensagem e segura quando o caso exige um humano. O fluxo abaixo é o coração do MVP.")
flow = doc.add_table(rows=1, cols=3); fixed_table(flow,[1.05,2.15,3.30])
for c, text in zip(flow.rows[0].cells,["ETAPA","O QUE ACONTECE","VALOR PARA A IMOBILIÁRIA"]): shade(c,LIGHT); set_cell_text(c,text,True,BLUE,9)
rows = [
    ("1", "Lead chama no WhatsApp", "Resposta imediata, mesmo fora do horário comercial."),
    ("2", "Agente identifica a intenção", "Diferencia compra, aluguel, dúvida inicial ou pedido de corretor."),
    ("3", "Conversa de qualificação", "Coleta dados essenciais sem parecer um formulário frio."),
    ("4", "Busca no catálogo", "Filtra por região, finalidade, tipo e orçamento; depois evolui para busca semântica."),
    ("5", "Apresenta opções e próximo passo", "Oferece até três alternativas e direciona para visita ou atendimento."),
    ("6", "Handoff ao corretor", "Entrega o contexto completo no Chatwoot para evitar repetição de perguntas."),
]
for n, action, value in rows:
    cells=flow.add_row().cells
    for cell, text in zip(cells,[n,action,value]): set_cell_text(cell,text,cell==cells[0],NAVY if cell!=cells[0] else GOLD)

doc.add_page_break()
add_heading(doc, "3. A experiência do usuário em uma conversa real", 1)
add_body(doc, "O agente não deve soar como um robô que despeja perguntas. Ele conduz uma conversa curta, clara e útil. Exemplo de jornada:")
journey = doc.add_table(rows=1, cols=2); fixed_table(journey,[1.75,4.75])
for c,text in zip(journey.rows[0].cells,["MOMENTO","EXPERIÊNCIA"]): shade(c,PALE); set_cell_text(c,text,True,BLUE,9)
examples = [
    ("Abertura", "\"Olá, Maria! Posso te ajudar a encontrar um imóvel. Você procura compra ou aluguel?\""),
    ("Descoberta", "O agente entende região, orçamento, quartos, prazo e necessidade de financiamento, adaptando perguntas ao que a pessoa já contou."),
    ("Recomendação", "\"Com base no que você busca, separei três opções em Moema dentro da sua faixa. Quer receber os detalhes?\""),
    ("Conversão", "Quando há interesse ou intenção forte, o agente propõe uma visita ou chama um corretor com o resumo do perfil."),
    ("Humano", "O corretor recebe: objetivo, preferências, orçamento, urgência, imóveis sugeridos e toda a conversa anterior.")
]
for a,b in examples:
    cells=journey.add_row().cells; set_cell_text(cells[0],a,True,DARK); set_cell_text(cells[1],b)

add_heading(doc, "Limites deliberados do agente", 2)
add_bullets(doc, [
    "Não inventa preço, disponibilidade, localização ou condição comercial.",
    "Não promete aprovação de financiamento, contratos ou condições jurídicas.",
    "Transfere para um humano quando recebe pedido explícito, baixa confiança ou caso fora da política.",
    "Registra conversas e deduplica reentregas do WhatsApp para preservar histórico confiável."
])

add_heading(doc, "4. Como vendemos", 1)
add_body(doc, "A venda inicial deve ser consultiva e focada em resultado operacional, não em ‘IA’. O comprador é o dono, diretor comercial ou gestor de atendimento que sente na prática a perda de leads e a sobrecarga dos corretores.")
add_heading(doc, "Cliente ideal para o piloto", 2)
add_bullets(doc, [
    "Imobiliária com volume recorrente de leads pelo WhatsApp, site, portais ou anúncios.",
    "Equipe comercial que já usa WhatsApp todos os dias e aceita centralizar o handoff em uma inbox.",
    "Catálogo minimamente organizado e alguém capaz de manter preço/disponibilidade atualizados.",
    "Decisor próximo da operação, disposto a acompanhar métricas por 30 a 60 dias."
])
add_heading(doc, "Oferta de entrada", 2)
add_numbered(doc, [
    "Diagnóstico de uma semana: entendemos o funil, as perguntas, o catálogo e as regras de passagem para corretores.",
    "Piloto guiado de 30 dias: um número, uma equipe e um conjunto de casos de uso bem definidos.",
    "Reunião semanal de resultados: corrigimos tom de voz, qualificação e regras conforme as conversas reais.",
    "Conversão para assinatura: preço mensal pela operação + implantação/configuração inicial."
])
add_body(doc, "Hipótese comercial inicial: cobrar implantação pelo trabalho de configuração e uma mensalidade por unidade/caixa de entrada, com franquia de conversas e custo adicional por volume. O preço final deve nascer de entrevistas e pilotos; antes disso, não devemos fingir precisão.")

doc.add_page_break()
add_heading(doc, "5. Tecnologia: o que já temos e por que importa", 1)
tech = doc.add_table(rows=1, cols=3); fixed_table(tech,[1.35,2.15,3.0])
for c,text in zip(tech.rows[0].cells,["CAMADA","TECNOLOGIA","PAPEL NO MVP"]): shade(c,LIGHT); set_cell_text(c,text,True,BLUE,9)
stack = [
    ("Canal", "WhatsApp Business API (Meta ou 360dialog)", "Recebe e envia mensagens. A escolha final depende de custo, suporte e elegibilidade do cliente."),
    ("Orquestração", "n8n", "Conecta webhooks, IA, API, catálogo e Chatwoot; acelera o piloto sem construir toda integração do zero."),
    ("Produto", "Spring Boot / Java 21", "Centraliza regras de negócio, histórico, qualificação e endpoints estáveis para o restante da plataforma."),
    ("Dados", "PostgreSQL 16 + pgvector", "Guarda leads, conversas, imóveis e permite busca estruturada e, depois, semântica."),
    ("Atendimento humano", "Chatwoot", "Inbox do corretor para assumir conversas com contexto."),
    ("IA", "Modelo de linguagem + embeddings", "Entende mensagens, orienta perguntas e encontra imóveis relevantes com limites e regras."),
    ("Infra", "Docker + Caddy + Redis", "Ambiente reproduzível, proxy TLS e filas para serviços auxiliares.")
]
for x,y,z in stack:
    cells=tech.add_row().cells; set_cell_text(cells[0],x,True,DARK,9.4); set_cell_text(cells[1],y,False,NAVY,9.4); set_cell_text(cells[2],z,False,NAVY,9.4)

add_heading(doc, "Por que esta arquitetura é boa para agora", 2)
add_body(doc, "Ela separa o que muda frequentemente do que precisa ser confiável. O n8n permite ajustar o fluxo de atendimento rapidamente. A API e o banco preservam regras, dados e integrações centrais. Mais adiante, as configurações por cliente - perguntas, tom, critérios de lead quente e base de conhecimento - viram a ponte para outros segmentos.")

add_heading(doc, "6. Contras, riscos e como lidar com eles", 1)
risks = doc.add_table(rows=1, cols=3); fixed_table(risks,[1.85,2.4,2.25])
for c,text in zip(risks.rows[0].cells,["RISCO / CONTRA","IMPACTO","RESPOSTA PRÁTICA"]): shade(c,PALE); set_cell_text(c,text,True,BLUE,9)
items = [
    ("Dados de imóveis desatualizados", "Recomendação errada derruba confiança.", "Começar com importação simples e regra de só oferecer imóveis ativos; evoluir integração com ERP/CRM."),
    ("IA responde algo incorreto", "Risco de frustração ou promessa indevida.", "Usar contexto restrito, regras claras, logs, mensagens de segurança e handoff por baixa confiança."),
    ("Dependência de WhatsApp", "Políticas, templates e janela de 24h afetam a operação.", "Desenhar mensagens de opt-in, templates aprovados e processos compatíveis com a política do provedor."),
    ("Venda longa ou sem dono claro", "Pilotos não viram receita.", "Vender para quem sofre a dor operacional e medir resultado desde a primeira semana."),
    ("Escopo amplo demais", "Atraso e produto genérico antes de validar valor.", "Focar em imobiliárias, um fluxo e 1-3 pilotos; generalizar apenas padrões comprovados."),
    ("Privacidade e LGPD", "Risco legal e reputacional.", "Coletar somente o necessário, definir retenção, controlar acesso e formalizar papéis com cada cliente.")
]
for a,b,c in items:
    cells=risks.add_row().cells; set_cell_text(cells[0],a,True,"9B1C1C",9.3); set_cell_text(cells[1],b,False,NAVY,9.3); set_cell_text(cells[2],c,False,NAVY,9.3)

doc.add_page_break()
add_heading(doc, "7. O que é MVP - e o que não é", 1)
add_body(doc, "O MVP precisa provar uma coisa: conseguimos melhorar o atendimento de entrada sem aumentar trabalho para a equipe. Não precisa nascer como uma plataforma completa para todos os setores.")
comparison = doc.add_table(rows=1, cols=2); fixed_table(comparison,[3.25,3.25])
for c,text in zip(comparison.rows[0].cells,["ENTRA NO MVP","FICA PARA DEPOIS"]): shade(c,PALE); set_cell_text(c,text,True,BLUE,9)
left = ["WhatsApp conectado a um fluxo real", "Qualificação imobiliária configurável", "Catálogo de demonstração e filtros", "Sugestão de até três imóveis", "Handoff com histórico para corretor", "Métricas básicas do funil", "Piloto com cliente real"]
right = ["CRM completo próprio", "App mobile para corretores", "Marketplace de imóveis", "Suporte simultâneo a muitos nichos", "Automação total de agendamento", "Integrações profundas com todos os ERPs", "Painel analítico avançado"]
for a,b in zip(left,right):
    cells=comparison.add_row().cells; set_cell_text(cells[0],"• " + a,False,NAVY,10); set_cell_text(cells[1],"• " + b,False,MUTED,10)

add_heading(doc, "8. Roteiro de 90 dias", 1)
roadmap = doc.add_table(rows=1, cols=3); fixed_table(roadmap,[1.25,2.55,2.70])
for c,text in zip(roadmap.rows[0].cells,["FASE","ENTREGA","SINAL DE SUCESSO"]): shade(c,LIGHT); set_cell_text(c,text,True,BLUE,9)
for phase, delivery, signal in [
    ("Dias 1-30", "Subir stack, validar API, configurar um fluxo n8n e um catálogo de demonstração.", "Conversa completa simulada: mensagem → qualificação → imóvel → handoff."),
    ("Dias 31-60", "Rodar com 1 cliente piloto e ajustar tom, perguntas e regras a partir de conversas reais.", "Leads respondidos, qualificados e transferidos sem perda de contexto."),
    ("Dias 61-90", "Medir resultado, fechar plano comercial e repetir implementação com segundo/terceiro cliente.", "Evidência de valor e processo de implantação replicável.")
]:
    cells=roadmap.add_row().cells; set_cell_text(cells[0],phase,True,DARK,9.5); set_cell_text(cells[1],delivery,False,NAVY,9.5); set_cell_text(cells[2],signal,False,NAVY,9.5)

add_heading(doc, "Decisão que precisamos tomar como sócios", 2)
add_body(doc, "Nosso foco não é provar que conseguimos fazer um chatbot. É provar que conseguimos gerar mais conversas qualificadas e mais visitas para uma imobiliária. A recomendação é assumir esse recorte, escolher um primeiro piloto e organizar a execução em torno de métricas de negócio - não de funcionalidades.", "A recomendação é assumir esse recorte")
add_bullets(doc, ["Definir quem lidera produto/cliente, tecnologia e comercial.", "Escolher o perfil do primeiro cliente-piloto e iniciar entrevistas de descoberta.", "Fechar o escopo de 30 dias e os indicadores que justificarão a assinatura mensal."])

doc.core_properties.title = "Imob Agent - Proposta de MVP para Sócios"
doc.core_properties.subject = "Visão de produto, funcionamento, estratégia comercial e riscos"
doc.core_properties.author = "Imob Agent"
doc.save(OUT)
print(OUT)
